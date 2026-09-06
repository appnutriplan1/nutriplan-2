from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = r"public/plantilla-plan-nutriplan.docx"
GREEN = "1F4A3F"
CREAM = "F1EBE1"

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)

def add_label(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(GREEN)
    p.add_run(value)

def add_meal(doc, day, meal, name, ingredients, steps, tip="", conservation=""):
    first = len(doc.paragraphs)
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(10)
    heading.paragraph_format.space_after = Pt(5)
    run = heading.add_run(f"{day}  |  {meal}")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(GREEN)
    add_label(doc, "DÍA", day)
    add_label(doc, "COMIDA", meal)
    add_label(doc, "NOMBRE", name)
    add_label(doc, "PORCIONES", "1")
    p = doc.add_paragraph()
    p.add_run("INGREDIENTES:").bold = True
    for qty, food, measure in ingredients:
        doc.add_paragraph(f"- {qty} g | {food} | {measure}")
    p = doc.add_paragraph()
    p.add_run("PREPARACIÓN:").bold = True
    for i, step in enumerate(steps, 1):
        doc.add_paragraph(f"{i}. {step}")
    add_label(doc, "CONSEJO", tip)
    add_label(doc, "CONSERVACIÓN", conservation)
    block = doc.paragraphs[first:]
    for paragraph in block[:-1]:
        paragraph.paragraph_format.keep_with_next = True

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8); section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2); section.right_margin = Cm(2)
styles = doc.styles
styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(10.5)
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Plantilla de plan semanal NutriPlan")
intro = doc.add_paragraph("Complete esta plantilla sin cambiar los encabezados. Repita el bloque para cada comida y para cada día. Puede cambiar los ejemplos por los alimentos y cantidades de su paciente.")
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
table = doc.add_table(rows=1, cols=1)
cell = table.cell(0,0); shade(cell, CREAM); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
cell.paragraphs[0].add_run("Formato obligatorio del ingrediente: gramos | nombre del alimento | medida casera. Ejemplo: 120 g | Pollo | 1 filete.").bold = True

add_meal(doc,"Lunes","Desayuno","Avena con plátano y leche",[(40,"Avena","4 cucharadas"),(200,"Leche","1 taza"),(100,"Plátano","1 unidad pequeña")],["Cocinar la avena con la leche.","Servir con el plátano."],"Puede añadir canela.","Consumir recién preparado.")
add_meal(doc,"Lunes","Media mañana","Yogur con fresas",[(170,"Yogur natural","1 envase"),(100,"Fresa","1 taza")],["Lavar y picar las fresas.","Servir con el yogur."],"Elegir yogur sin azúcar.","Mantener refrigerado.")
add_meal(doc,"Lunes","Almuerzo","Pollo con camote y ensalada",[(120,"Pollo","1 filete"),(150,"Camote","1 unidad mediana"),(5,"Aceite de oliva","1 cucharadita")],["Cocinar el pollo a la plancha.","Hervir el camote y servir con ensalada."],"Puede reemplazar el camote por papa.","Refrigerar hasta 2 días.")
add_meal(doc,"Lunes","Media tarde","Pan integral con palta",[(60,"Pan integral","2 rebanadas"),(50,"Palta","1 cuarto de unidad")],["Tostar el pan.","Agregar la palta triturada."],"Añadir unas gotas de limón.","Preparar al momento.")
add_meal(doc,"Lunes","Cena","Tortilla de verduras",[(100,"Huevo","2 unidades"),(100,"Verduras","1 taza"),(5,"Aceite","1 cucharadita")],["Batir los huevos.","Agregar las verduras y cocinar a fuego medio."],"Usar las verduras disponibles.","Refrigerar hasta 1 día.")

doc.add_page_break()
h = doc.add_paragraph()
r = h.add_run("Cómo continuar con el martes")
r.bold = True; r.font.size = Pt(18); r.font.color.rgb = RGBColor.from_string(GREEN)
doc.add_paragraph("Después de las cinco comidas del lunes, copie el mismo bloque y cambie el nombre del día de Lunes a Martes. Continúe con las demás comidas del martes y repita el procedimiento hasta el domingo.")
add_meal(doc,"Martes","Desayuno","Escriba aquí el nombre",[(0,"Escriba aquí el alimento","Escriba aquí la medida")],["Escriba aquí el primer paso.","Escriba aquí el segundo paso."],"Escriba aquí el consejo opcional.","Escriba aquí la conservación opcional.")
doc.add_paragraph("Luego agregue: Media mañana, Almuerzo, Media tarde y Cena del martes usando el mismo formato.")

doc.save(OUT)
print(OUT)
